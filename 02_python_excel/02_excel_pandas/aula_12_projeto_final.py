import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

caminho_invoice = (
    "02_python_excel/02_excel_pandas/invoice_aula12.xlsx"
)

caminho_pl = (
     "02_python_excel/02_excel_pandas/packing_list_aula12.xlsx"
)

invoice = pd.read_excel(
    caminho_invoice,
    dtype={
        "codigo": str,
        "ncm": str
        }
)

packing_list = pd.read_excel(
    caminho_pl,
    dtype={
        "codigo": str,
        }
)

invoice["codigo"] = invoice["codigo"].str.zfill(3)
invoice["valor_unitario"] = (
    invoice["valor_unitario"]
    .str.replace(",", ".")
    .astype(float)
)
packing_list["codigo"] = packing_list["codigo"].str.zfill(3)

print("INVOICE")
print(invoice)

print()

print("PACKING LIST")
print(packing_list)

duplicados_invoice = invoice[
    invoice["codigo"].duplicated(keep=False)
]

duplicados_pl = packing_list[
    packing_list["codigo"].duplicated(keep=False)
]

print()
print("DUPLICADO NA INVOICE")
print(duplicados_invoice)

print()
print("DUPLICADO NO PACKING LIST")
print(duplicados_pl)

if not duplicados_invoice.empty:
    raise ValueError(
        "Exitem dódigos duplicados na Invoice. "
        "Corrija a base antes de continuar."
    )

if not duplicados_pl.empty:
    raise ValueError(
        "Existem códigos duplicados no Packing List. "
        "Corrija a base antes de continuar."
    )

conferencia = pd.merge(
    invoice,
    packing_list,
    on="codigo",
    how="outer",
    indicator=True
)

conferencia["status"] = "OK"

conferencia.loc[
    conferencia["_merge"] == "left_only",
    "status"
] = "AUSENTE NO PL"

conferencia.loc[
    conferencia["_merge"] == "right_only",
    "status"
] = "AUSENTE NA INVOICE"

conferencia.loc[
    (
        conferencia["quantidade_invoice"].isna()
        | conferencia["quantidade_pl"].isna()
    )
    & (conferencia["_merge"] == "both"),
    "status"
] = "DADOS INCOMPLETOS"

conferencia.loc[
    (conferencia["_merge"] == "both")
    & (conferencia["quantidade_invoice"].notna()
    & conferencia["quantidade_pl"].notna()
    & (
        conferencia["quantidade_invoice"]
        != conferencia["quantidade_pl"]
    )
    ),
    "status"
] = "DIVERGÊNCIA"

conferencia["status_peso"] = "OK"

conferencia.loc[
    conferencia["peso_bruto"] < conferencia["peso_liquido"],
    "status_peso"
] = "ERRO DE PESO"

conferencia.loc[
    conferencia["peso_liquido"].isna() | conferencia["peso_bruto"].isna(),
    "status_peso"
] = "NAO APLICAVEL"

conferencia["diferenca"] = (
    conferencia["quantidade_invoice"] - conferencia["quantidade_pl"]
).abs()

conferencia["valor_total"] = (
    conferencia["quantidade_invoice"] 
    * conferencia["valor_unitario"]
)

conferencia["peso_liquido_unitario"] = (
    conferencia["peso_liquido"] 
    / conferencia["quantidade_pl"]
)

conferencia["peso_bruto_unitario"] = (
    conferencia["peso_bruto"]
    / conferencia["quantidade_pl"]
)

print()
print("VALIDAÇÃO DOS CÁLCULOS")

print(
    conferencia[
        [
            "codigo",
            "diferenca",
            "valor_total",
            "peso_liquido_unitario",
            "peso_bruto_unitario"
        ]
    ]
)

conferencia["status_geral"] = conferencia["status"]

conferencia.loc[
    (conferencia["status"] == "OK")
    & (conferencia["status_peso"] == "ERRO DE PESO"),
    "status_geral"
] = "ERRO DE PESO"

conferencia.loc[
    (conferencia["status"] == "DIVERGÊNCIA")
    & (conferencia["status_peso"] == "ERRO DE PESO"),
    "status_geral"
] = "REVISAR"

print(
    conferencia[
        [
            "codigo",
            "status",
            "status_peso",
            "status_geral"
        ]
    ]
)

print()
print("CONFERÊNCIA FINAL")
print(conferencia)

caminho_saida = (
    "02_python_excel/02_excel_pandas/"
    "relatorio_conferencia_aula12.xlsx"
)

conferencia.to_excel(
    caminho_saida,
    index=False
)

workbook = load_workbook(caminho_saida)
planilha = workbook.active

for cell in planilha[1]:
    cell.font = Font(bold=True)

planilha.freeze_panes = "A2"

planilha.auto_filter.ref = planilha.dimensions

for coluna in planilha.columns:
    maior_tamanho = 0

    for celula in coluna:
        if celula.value is not None:
            tamanho = len(str(celula.value))

            if tamanho > maior_tamanho:
                maior_tamanho = tamanho

    letra_coluna = get_column_letter(coluna[0].column)

    planilha.column_dimensions[letra_coluna].width = (
        maior_tamanho + 2
    )

cores_status_geral = {
    "OK": "C6EFCE",
    "DIVERGÊNCIA": "FFEB9C",
    "ERRO DE PESO": "FFEB9C",
    "DADOS INCOMPLETOS": "FFC000",
    "AUSENTE NA INVOICE": "FFC7CE",
    "AUSENTE NO PL": "FFC7CE",
    "REVISAR": "CDA4DE",
}

colunas = {
    cell.value: cell.column
    for cell in planilha[1]
}

coluna_status_geral = colunas["status_geral"]

for linha in range(2, planilha.max_row + 1):
    valor_status_geral = planilha.cell(
        row=linha,
        column=coluna_status_geral
    ).value

    cor = cores_status_geral.get(valor_status_geral)

    if cor:
        preenchimento = PatternFill(
            start_color=cor,
            end_color=cor,
            fill_type="solid"
        )

        for coluna in range(1, planilha.max_column + 1):
            planilha.cell(
                row=linha,
                column=coluna
            ).fill = preenchimento

workbook.save(caminho_saida)

print()
print("Relatório da Aula 12 criado com sucesso.")