from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

caminho = (
    "02_python_excel/02_excel_pandas/"
    "resultado_conferencia_aula10.xlsx"
)

workbook = load_workbook(caminho)
planilha = workbook.active

for celula in planilha[1]: 
    celula.font = Font(bold=True)

planilha.freeze_panes = "A2"

planilha.auto_filter.ref = planilha.dimensions

for coluna in planilha.columns:
    maior_tamanho = 0

    for celula in coluna:
        if celula.value is not None:
            tamanho = len(str(celula.value))

            if tamanho > maior_tamanho:
                maior_tamanho = tamanho

    letra_coluna = get_column_letter(coluna[0].column)
    planilha.column_dimensions[letra_coluna].width = maior_tamanho + 2

coluna_status = None
coluna_status_peso = None

for celula in planilha[1]:

    if celula.value == "status":
        coluna_status = celula.column

    elif celula.value == "status_peso":
        coluna_status_peso = celula.column

if coluna_status is None or coluna_status_peso is None:
    raise ValueError(
        "Coluna 'status' ou 'status_peso' - não encontradas"
    )

print("Coluna status:", coluna_status)
print("Coluna status_peso:", coluna_status_peso)


coluna_status_geral = None

for celula in planilha[1]:
    if celula.value == "status_geral":
        coluna_status_geral = celula.column
        break

if coluna_status_geral is None:
    raise ValueError(
    "Coluna 'status_geral' não encontrada."
    )

preenchimento_ok = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

preenchimento_divergencia = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC"
)

preenchimento_erro = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC"
)

for linha in range(2, planilha.max_row + 1):

    status = planilha.cell(
        row=linha,
        column=coluna_status
    ).value

    status_peso = planilha.cell(
        row=linha,
        column=coluna_status_peso
    ).value

    print(
        "Linha", linha,
        "| status:", status,
        "| status_peso:", status_peso
    )

for linha in range(2, planilha.max_row + 1):

    status_geral = planilha.cell(
        row=linha,
        column=coluna_status_geral
    ).value

    if status_geral == "OK":
        preenchimento = preenchimento_ok

    elif status_geral == "DIVERGÊNCIA":
        preenchimento = preenchimento_divergencia

    elif status_geral in ["ERRO DE PESO", "REVISAR"]:
        preenchimento = preenchimento_erro

    else:
        continue

    for coluna in range(1, planilha.max_column + 1):
        planilha.cell(
            row=linha,
            column=coluna
        ).fill = preenchimento

workbook.save(
     "02_python_excel/02_excel_pandas/"
    "resultado_conferencia_aula10_formatado.xlsx"
)

print("Relatório da Aula 10 formatado com sucesso.")