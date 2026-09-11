from openpyxl import load_workbook

caminho = (
    "02_python_excel/02_excel_pandas/"
    "resultado_conferencia_aula08.xlsx"
)

workbook = load_workbook(caminho)
planilha = workbook.active

from openpyxl.styles import PatternFill, Font

preenchimento_ok = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

preenchimento_divergencia = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC"
)

preenchimento_incompleto = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6"
)
preenchimento_ausente = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC"
)

coluna_status = None
for celula in planilha[1]:
    if celula.value == "status":
        coluna_status = celula.column
        break

for linha in range(2,planilha.max_row + 1):
    status = planilha.cell(
        row=linha,
        column=coluna_status
    ).value

    if status == "OK":
        preenchimento = preenchimento_ok

    elif status == "DIVERGÊNCIA":
        preenchimento = preenchimento_divergencia

    elif status == "DADOS INCOMPLETOS":
        preenchimento = preenchimento_incompleto

    elif status in ["AUSENTE NO PL", "AUSENTE NA INVOICE"]:
        preenchimento = preenchimento_ausente

    else:
        continue
    for coluna in range(1, planilha.max_column + 1):
        planilha.cell(
            row=linha,
            column=coluna
        ).fill = preenchimento

from openpyxl.utils import get_column_letter

for coluna in planilha.columns:
    maior_tamanho = 0

    for celula in coluna:
        if celula.value is not None:
            tamanho = len(str(celula.value))

            if tamanho > maior_tamanho:
                maior_tamanho = tamanho

    letra_coluna = get_column_letter(coluna[0].column)

    planilha.column_dimensions[letra_coluna].width = maior_tamanho + 2

    print(
    "Coluna:",
    letra_coluna,
    "| Maior tamanho:",
    maior_tamanho,
    "| Largura definida:",
    maior_tamanho + 2
)

for celula in planilha[1]:
    celula.font = Font(bold=True)
    planilha.freeze_panes = "A2"

planilha.auto_filter.ref = planilha.dimensions

workbook.save(
    "02_python_excel/02_excel_pandas/"
    "resultado_conferencia_formatado.xlsx"
)

print("Arquivo formatado criado com sucesso.")